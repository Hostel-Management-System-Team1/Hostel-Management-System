
import csv
from datetime import date, timedelta
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from django_otp.plugins.otp_totp.models import TOTPDevice
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.complaints.models import Complaint
from apps.mess.models import MessActivity


class ReportingAccessTests(TestCase):
    def setUp(self):
        self.warden = User.objects.create_user(
            username="report_warden",
            password="TestPassword123!",
            role=User.Role.WARDEN,
        )
        self.resident = User.objects.create_user(
            username="report_resident",
            password="TestPassword123!",
            role=User.Role.RESIDENT,
        )

    def log_in_verified_warden(self, user):
        """Log in a test Warden with a verified OTP device."""
        device = TOTPDevice.objects.create(
            user=user,
            name="Test authenticator",
            confirmed=True,
        )
        self.client.force_login(user)

        session = self.client.session
        session["otp_device_id"] = device.persistent_id
        session.save()

    def test_anonymous_user_is_redirected_to_login(self):
        response = self.client.get(
            reverse("reporting:dashboard")
        )

        self.assertEqual(response.status_code, 302)
        self.assertIn("/accounts/login/", response.url)

    def test_warden_can_open_reports_dashboard(self):
        self.log_in_verified_warden(self.warden)

        response = self.client.get(
            reverse("reporting:dashboard")
        )

        self.assertEqual(response.status_code, 200)

    def test_resident_is_forbidden_from_reports(self):
        self.client.force_login(self.resident)

        response = self.client.get(
            reverse("reporting:dashboard")
        )

        self.assertEqual(response.status_code, 403)


class ReportingExportTests(TestCase):
    def setUp(self):
        self.warden = User.objects.create_user(
            username="export_warden",
            password="TestPassword123!",
            role=User.Role.WARDEN,
        )

        self.device = TOTPDevice.objects.create(
            user=self.warden,
            name="Test authenticator",
            confirmed=True,
        )

        self.client.force_login(self.warden)

        session = self.client.session
        session["otp_device_id"] = self.device.persistent_id
        session.save()

        self.complaint = Complaint.objects.create(
            title="Leaking bathroom tap",
            description="The bathroom tap needs repair.",
            category=Complaint.Category.MAINTENANCE,
            status=Complaint.Status.OPEN,
            created_by=self.warden,
            created_at=timezone.now() - timedelta(days=4),
        )

        self.mess_activity = MessActivity.objects.create(
            activity_date=date.today(),
            meal_type=MessActivity.MealType.LUNCH,
            meals_served=120,
            notes="Lunch service test record",
            recorded_by=self.warden,
        )

    def test_occupancy_export_returns_csv(self):
        response = self.client.get(
            reverse("reporting:export_occupancy")
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertIn(
            "attachment; filename=",
            response["Content-Disposition"],
        )

        rows = list(
            csv.reader(
                response.content.decode("utf-8").splitlines()
            )
        )

        self.assertEqual(rows[0][0], "Block")
        self.assertIn("Occupied beds", rows[0])

    def test_dues_export_returns_csv(self):
        response = self.client.get(
            reverse("reporting:export_dues")
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])

        rows = list(
            csv.reader(
                response.content.decode("utf-8").splitlines()
            )
        )

        self.assertEqual(rows[0][0], "Invoice")
        self.assertIn("Amount", rows[0])

    def test_occupancy_export_returns_xlsx(self):
        response = self.client.get(
            reverse("reporting:export_occupancy_xlsx")
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn(".xlsx", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"PK"))

    def test_dues_export_returns_xlsx(self):
        response = self.client.get(
            reverse("reporting:export_dues_xlsx")
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn(".xlsx", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"PK"))

    def test_occupancy_export_returns_pdf(self):
        response = self.client.get(
            reverse("reporting:export_occupancy_pdf")
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/pdf",
        )
        self.assertIn(".pdf", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_dues_export_returns_pdf(self):
        response = self.client.get(
            reverse("reporting:export_dues_pdf")
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/pdf",
        )
        self.assertIn(".pdf", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_complaints_csv_contains_real_records(self):
        response = self.client.get(
            reverse("reporting:export_complaints")
        )

        self.assertEqual(response.status_code, 200)

        rows = list(
            csv.reader(
                response.content.decode("utf-8").splitlines()
            )
        )

        self.assertEqual(rows[0][0], "Title")
        self.assertTrue(
            any("Leaking bathroom tap" in row for row in rows)
        )
        self.assertTrue(
            any(row[-1] == "4" for row in rows[1:])
        )

    def test_mess_csv_contains_real_records(self):
        response = self.client.get(
            reverse("reporting:export_mess")
        )

        self.assertEqual(response.status_code, 200)

        rows = list(
            csv.reader(
                response.content.decode("utf-8").splitlines()
            )
        )

        self.assertEqual(rows[0][0], "Activity date")
        self.assertTrue(
            any("120" in row for row in rows)
        )
        self.assertTrue(
            any("Lunch" in row for row in rows)
        )

    def test_complaints_xlsx_contains_real_records(self):
        response = self.client.get(
            reverse("reporting:export_complaints_xlsx")
        )

        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(
            BytesIO(response.content),
            read_only=True,
        )
        try:
            rows = list(workbook.active.values)
        finally:
            workbook.close()

        self.assertEqual(rows[0][0], "Title")
        self.assertTrue(
            any(
                "Leaking bathroom tap" in row
                for row in rows
            )
        )

    def test_mess_xlsx_contains_real_records(self):
        response = self.client.get(
            reverse("reporting:export_mess_xlsx")
        )

        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(
            BytesIO(response.content),
            read_only=True,
        )
        try:
            rows = list(workbook.active.values)
        finally:
            workbook.close()

        self.assertEqual(rows[0][0], "Activity date")
        self.assertTrue(
            any(120 in row for row in rows)
        )

    def test_complaints_pdf_returns_valid_pdf(self):
        response = self.client.get(
            reverse("reporting:export_complaints_pdf")
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/pdf",
        )
        self.assertIn(".pdf", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_mess_pdf_returns_valid_pdf(self):
        response = self.client.get(
            reverse("reporting:export_mess_pdf")
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/pdf",
        )
        self.assertIn(".pdf", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))
